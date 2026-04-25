#!/usr/bin/env python3
"""Generate print-ready sticker artwork for 3DPawsnToys pet toys.

Produces in stickers/:
  sticker_single.pdf      - single 80x50mm sticker, exact size (label printers).
  sticker_sheet_a4.pdf    - A4 page with 8 stickers + cut guides (home printers).
  sticker_template.docx   - editable Word/Pages version (single sticker).
  sticker_preview.html    - browser preview at exact size.
"""
from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor, black, grey
from reportlab.pdfgen import canvas
from reportlab.pdfbase.pdfmetrics import stringWidth

from docx import Document
from docx.shared import Mm, Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
QR_PATH = ROOT / "qr" / "qr_code.png"
LOGO_PATH = ROOT / "assets" / "logo.png"
OUT = ROOT / "stickers"
OUT.mkdir(exist_ok=True)

URL = "samshahdad.github.io/pet-toy-safety"
BRAND = "3DPawsnToys"
EMAIL = "contact@3dpawsandtoys.com"

ORANGE = HexColor("#ef6a2a")
BLUE = HexColor("#1f4ea1")
DARK = HexColor("#1a1a1a")

STICKER_W = 80 * mm
STICKER_H = 50 * mm

DISCLAIMER_LINES = [
    ("title",   "SAFETY WARNING - PLEASE READ"),
    ("body",    "3D-printed pet toy made from PLA / PETG / TPU. Non-toxic, NOT edible."),
    ("body",    "For dogs & cats only - always supervise play."),
    ("body",    "Designed for light-to-moderate chewers; not for power chewers."),
    ("body",    "Choose a size LARGER than your pet's mouth to reduce choking risk."),
    ("body",    "Inspect before & after each use. Discard if cracked, broken, or chewed"),
    ("body",    "into pieces - fragments may cause choking, dental injury, or blockage."),
    ("body",    "Not for children under 3. Clean with mild soap & water; keep from heat."),
    ("body",    "Used at owner's risk. Manufacturer not liable for injury from misuse."),
]


def draw_sticker(c: canvas.Canvas, x: float, y: float, w: float = STICKER_W, h: float = STICKER_H,
                 outline: bool = False):
    """Draw one sticker with bottom-left at (x, y)."""

    if outline:
        c.setStrokeColor(grey)
        c.setLineWidth(0.25)
        c.setDash(2, 2)
        c.rect(x, y, w, h)
        c.setDash()

    pad = 2.5 * mm
    qr_size = 28 * mm
    qr_x = x + pad
    qr_y = y + (h - qr_size) / 2
    c.drawImage(str(QR_PATH), qr_x, qr_y, qr_size, qr_size,
                preserveAspectRatio=True, mask='auto')

    c.setFillColor(grey)
    c.setFont("Helvetica", 4.5)
    scan_text = "Scan for full safety info"
    tw = stringWidth(scan_text, "Helvetica", 4.5)
    c.drawString(qr_x + (qr_size - tw) / 2, qr_y - 1.6 * mm, scan_text)

    text_x = x + pad + qr_size + 2 * mm
    text_w = w - (text_x - x) - pad
    text_top = y + h - pad

    c.setFillColor(BLUE)
    c.setFont("Helvetica-Bold", 7)
    c.drawString(text_x, text_top - 6, BRAND)

    c.setFillColor(ORANGE)
    c.setFont("Helvetica-Bold", 5.5)
    title_y = text_top - 11
    c.drawString(text_x, title_y, "\u26A0 " + DISCLAIMER_LINES[0][1])

    c.setFillColor(DARK)
    line_y = title_y - 6
    line_height = 5.6
    c.setFont("Helvetica", 4.6)
    for kind, line in DISCLAIMER_LINES[1:]:
        c.drawString(text_x, line_y, line)
        line_y -= line_height

    c.setFillColor(BLUE)
    c.setFont("Helvetica-Oblique", 4.6)
    c.drawString(text_x, y + pad + 1, URL + "  |  " + EMAIL)


def make_single_pdf():
    out = OUT / "sticker_single.pdf"
    c = canvas.Canvas(str(out), pagesize=(STICKER_W, STICKER_H))
    c.setTitle("3DPawsnToys Safety Sticker - 80x50mm")
    draw_sticker(c, 0, 0, outline=False)
    c.showPage()
    c.save()
    print("Wrote", out)


def make_sheet_pdf():
    """A4 sheet with 4 columns x ... no - 2 columns x 4 rows = 8 stickers. With cut marks."""
    out = OUT / "sticker_sheet_a4.pdf"
    c = canvas.Canvas(str(out), pagesize=A4)
    c.setTitle("3DPawsnToys Safety Stickers - A4 sheet, 8 per page")

    page_w, page_h = A4
    cols, rows = 2, 4

    grid_w = cols * STICKER_W
    grid_h = rows * STICKER_H
    margin_x = (page_w - grid_w) / 2
    margin_y = (page_h - grid_h) / 2

    for r in range(rows):
        for col in range(cols):
            x = margin_x + col * STICKER_W
            y = margin_y + r * STICKER_H
            draw_sticker(c, x, y, outline=True)

    c.setStrokeColor(black)
    c.setLineWidth(0.25)
    mark_len = 3 * mm
    for r in range(rows + 1):
        y = margin_y + r * STICKER_H
        c.line(margin_x - mark_len - 1, y, margin_x - 1, y)
        c.line(margin_x + grid_w + 1, y, margin_x + grid_w + 1 + mark_len, y)
    for col in range(cols + 1):
        x = margin_x + col * STICKER_W
        c.line(x, margin_y - mark_len - 1, x, margin_y - 1)
        c.line(x, margin_y + grid_h + 1, x, margin_y + grid_h + 1 + mark_len)

    c.setFillColor(grey)
    c.setFont("Helvetica", 7)
    c.drawString(margin_x, margin_y + grid_h + 6 * mm,
                 "3DPawsnToys safety stickers - 80mm x 50mm each. Cut along dashed lines (use marks as guides).")
    c.drawString(margin_x, margin_y - 6 * mm,
                 "Print at 100% / 'Actual size' - do NOT 'Fit to page'. Test-scan the QR before printing the full batch.")

    c.showPage()
    c.save()
    print("Wrote", out)


def _set_cell_borders(cell, color="BFBFBF", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "dashed")
        b.set(qn("w:sz"), size)
        b.set(qn("w:color"), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def _set_cell_width(cell, width_mm):
    cell.width = Mm(width_mm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(int(width_mm * 56.7)))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def make_docx():
    out = OUT / "sticker_template.docx"
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)

    intro = doc.add_paragraph()
    run = intro.add_run("3DPawsnToys Safety Sticker (editable)")
    run.bold = True
    run.font.size = Pt(11)
    intro2 = doc.add_paragraph()
    intro2.add_run(
        "Sticker size: 80 mm wide x 50 mm tall. To print, set page margins to "
        "match or use the supplied PDF. The dashed border shows the cut line; "
        "remove or hide it before final printing."
    ).font.size = Pt(9)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    row = table.rows[0]
    row.height = Mm(50)

    qr_cell, text_cell = row.cells
    _set_cell_width(qr_cell, 32)
    _set_cell_width(text_cell, 48)
    _set_cell_borders(qr_cell)
    _set_cell_borders(text_cell)

    qr_para = qr_cell.paragraphs[0]
    qr_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr_run = qr_para.add_run()
    qr_run.add_picture(str(QR_PATH), width=Mm(28))
    cap = qr_cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run("Scan for full safety info")
    cap_run.font.size = Pt(5)
    cap_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    p_brand = text_cell.paragraphs[0]
    r_brand = p_brand.add_run(BRAND)
    r_brand.bold = True
    r_brand.font.size = Pt(8)
    r_brand.font.color.rgb = RGBColor(0x1F, 0x4E, 0xA1)

    p_warn = text_cell.add_paragraph()
    r_warn = p_warn.add_run("\u26A0 SAFETY WARNING - PLEASE READ")
    r_warn.bold = True
    r_warn.font.size = Pt(6)
    r_warn.font.color.rgb = RGBColor(0xEF, 0x6A, 0x2A)

    body_text = (
        "3D-printed pet toy made from PLA / PETG / TPU. Non-toxic, NOT edible. "
        "For dogs & cats only - always supervise play. "
        "Designed for light-to-moderate chewers; not for power chewers. "
        "Choose a size LARGER than your pet's mouth to reduce choking risk. "
        "Inspect before & after each use. Discard if cracked, broken, or chewed "
        "into pieces - fragments may cause choking, dental injury, or blockage. "
        "Not for children under 3. Clean with mild soap & water; keep from heat. "
        "Used at owner's risk. Manufacturer not liable for injury from misuse."
    )
    p_body = text_cell.add_paragraph()
    r_body = p_body.add_run(body_text)
    r_body.font.size = Pt(5)

    p_url = text_cell.add_paragraph()
    r_url = p_url.add_run(f"{URL}  |  {EMAIL}")
    r_url.font.size = Pt(5)
    r_url.italic = True
    r_url.font.color.rgb = RGBColor(0x1F, 0x4E, 0xA1)

    doc.add_paragraph()
    notes = doc.add_paragraph()
    notes.add_run(
        "Print tip: when exporting / printing, choose 'Actual size' or 100% scaling, "
        "not 'Fit to page'. Always test-scan the QR before printing a full batch."
    ).font.size = Pt(8)

    doc.save(str(out))
    print("Wrote", out)


def make_html_preview():
    """Browser preview at exact size."""
    body_text = (
        "3D-printed pet toy made from PLA / PETG / TPU. Non-toxic, NOT edible. "
        "For dogs &amp; cats only &mdash; always supervise play. "
        "Designed for light-to-moderate chewers; not for power chewers. "
        "Choose a size LARGER than your pet's mouth to reduce choking risk. "
        "Inspect before &amp; after each use. Discard if cracked, broken, or chewed "
        "into pieces &mdash; fragments may cause choking, dental injury, or blockage. "
        "Not for children under 3. Clean with mild soap &amp; water; keep from heat. "
        "Used at owner's risk. Manufacturer not liable for injury from misuse."
    )
    html = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8" />
<title>3DPawsnToys Sticker Preview - 80x50mm</title>
<style>
  @page {{ size: 80mm 50mm; margin: 0; }}
  body {{ background: #ddd; margin: 0; padding: 24px; font-family: Helvetica, Arial, sans-serif; }}
  .sticker {{
    width: 80mm; height: 50mm; box-sizing: border-box;
    background: #fff; padding: 2.5mm; display: flex; gap: 2mm;
    border: 0.25mm dashed #bbb; margin: 0 auto;
  }}
  .qr {{ width: 28mm; flex-shrink: 0; display: flex; flex-direction: column; align-items: center; justify-content: center; }}
  .qr img {{ width: 28mm; height: 28mm; }}
  .qr .cap {{ font-size: 4.5pt; color: #666; margin-top: 0.5mm; text-align: center; }}
  .body {{ flex: 1; display: flex; flex-direction: column; min-width: 0; }}
  .brand {{ font-size: 7.5pt; font-weight: 700; color: #1f4ea1; line-height: 1; }}
  .warn {{ font-size: 6pt; font-weight: 700; color: #ef6a2a; margin-top: 0.6mm; line-height: 1.1; }}
  .text {{ font-size: 4.6pt; color: #1a1a1a; margin-top: 0.5mm; line-height: 1.2; }}
  .url {{ font-size: 4.6pt; font-style: italic; color: #1f4ea1; margin-top: auto; }}
  @media print {{ body {{ background: #fff; padding: 0; }} .sticker {{ border: none; }} }}
</style>
</head>
<body>
  <div class="sticker">
    <div class="qr">
      <img src="../qr/qr_code.png" alt="QR" />
      <div class="cap">Scan for full safety info</div>
    </div>
    <div class="body">
      <div class="brand">{BRAND}</div>
      <div class="warn">&#9888; SAFETY WARNING &mdash; PLEASE READ</div>
      <div class="text">{body_text}</div>
      <div class="url">{URL} &nbsp;|&nbsp; {EMAIL}</div>
    </div>
  </div>
</body>
</html>"""
    out = OUT / "sticker_preview.html"
    out.write_text(html, encoding="utf-8")
    print("Wrote", out)


if __name__ == "__main__":
    make_single_pdf()
    make_sheet_pdf()
    make_docx()
    make_html_preview()
    print("\nAll sticker files written to:", OUT)
